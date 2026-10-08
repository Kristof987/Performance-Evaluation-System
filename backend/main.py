from copy import deepcopy
from datetime import date
import os
from uuid import uuid4

from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pydantic import TypeAdapter, ValidationError
from sqlalchemy import func
from sqlalchemy.orm import Session

from database import engine, Base, SessionLocal
import models
import schemas

#Base.metadata.create_all(bind=engine)

app = FastAPI()

FRONTEND_URL = os.getenv("FRONTEND_URL")

allowed_origins = [
    "http://localhost:5173",
    "http://127.0.0.1:5173",
]

if FRONTEND_URL:
    allowed_origins.append(FRONTEND_URL)

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

email_validator = TypeAdapter(schemas.EmailStr)
EMPLOYEE_IMPORT_HEADERS = [
    "Employee Name",
    "Employee Email Address",
    "Employee Company Role",
    "Employee System Role",
]
GROUP_IMPORT_HEADERS = ["Group Name", "Group Description"]


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@app.post("/users", response_model=schemas.UserResponse)
def create_user(user: schemas.UserCreate, db: Session = Depends(get_db)):
    new_user = models.User(
        username=user.username,
        email=user.email,
        password_hash=user.password,  # ezt mindjárt javítjuk hashelésre
        company_role_id=user.company_role_id,
        system_role_id=user.system_role_id,
    )

    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    return new_user


@app.post("/login", response_model=schemas.UserResponse)
def login(user_login: schemas.UserLogin, db: Session = Depends(get_db)):
    result = (
        db.query(models.User, models.CompanyRole.name.label("role_name"))
        .join(models.CompanyRole, models.CompanyRole.id == models.User.company_role_id)
        .filter(models.User.username == user_login.username.strip())
        .filter(models.User.is_active.is_(True))
        .first()
    )

    if result is None:
        raise HTTPException(status_code=401, detail="Nem letezo vagy inaktiv user")

    user, role_name = result

    user.last_login = func.now()
    db.commit()
    db.refresh(user)

    return {
        **user.__dict__,
        "role_name": role_name,
    }


@app.get("/people", response_model=schemas.PeopleResponse)
def get_people(db: Session = Depends(get_db)):
    users = (
        db.query(models.User)
        .order_by(models.User.is_active.desc(), models.User.username.asc())
        .all()
    )
    roles_by_id = {
        role.id: role.name
        for role in db.query(models.CompanyRole).all()
    }
    groups = db.query(models.CompanyGroup).order_by(models.CompanyGroup.name.asc()).all()
    groups_by_id = {group.id: group for group in groups}
    group_member_counts = {group.id: 0 for group in groups}
    group_ids_by_user_id = {user.id: [] for user in users}

    memberships = db.query(
        models.user_company_groups.c.user_id,
        models.user_company_groups.c.company_group_id,
    ).all()
    for membership in memberships:
        group = groups_by_id.get(membership.company_group_id)
        if group is None:
            continue
        group_member_counts[group.id] = group_member_counts.get(group.id, 0) + 1
        if membership.user_id in group_ids_by_user_id:
            group_ids_by_user_id[membership.user_id].append(group.id)

    def serialize_group(group: models.CompanyGroup):
        return {
            "id": group.id,
            "name": group.name,
            "description": group.description,
            "member_count": group_member_counts.get(group.id, 0),
        }

    return {
        "employees": [
            {
                "id": user.id,
                "username": user.username,
                "email": user.email,
                "company_role_id": user.company_role_id,
                "role_name": roles_by_id.get(user.company_role_id, "Unknown role"),
                "groups": [
                    serialize_group(groups_by_id[group_id])
                    for group_id in group_ids_by_user_id.get(user.id, [])
                    if group_id in groups_by_id
                ],
                "is_active": user.is_active,
            }
            for user in users
        ],
        "groups": [serialize_group(group) for group in groups],
    }


def get_or_create_company_role(role_name: str, db: Session):
    company_role = (
        db.query(models.CompanyRole)
        .filter(func.lower(models.CompanyRole.name) == role_name.lower())
        .first()
    )
    if company_role is None:
        company_role = models.CompanyRole(name=role_name)
        db.add(company_role)
        db.flush()
    return company_role


def get_or_create_system_role(role_name: str, db: Session):
    system_role = (
        db.query(models.SystemRole)
        .filter(func.lower(models.SystemRole.name) == role_name.lower())
        .first()
    )
    if system_role is None:
        system_role = models.SystemRole(name=role_name)
        db.add(system_role)
        db.flush()
    return system_role


@app.post("/people/employees", response_model=schemas.PeopleEmployeeResponse)
def create_people_employee(employee: schemas.PeopleEmployeeCreate, db: Session = Depends(get_db)):
    name = employee.name.strip()
    role_name = employee.role.strip()
    if name == "" or role_name == "":
        raise HTTPException(status_code=400, detail="Name and role are required")

    existing_user = (
        db.query(models.User)
        .filter((models.User.username == name) | (models.User.email == employee.email))
        .first()
    )
    if existing_user is not None:
        raise HTTPException(status_code=400, detail="Employee name or email already exists")

    company_role = get_or_create_company_role(role_name, db)
    system_role = get_or_create_system_role("Employee", db)

    group_ids = sorted(set(employee.group_ids))
    groups = []
    if len(group_ids) > 0:
        groups = (
            db.query(models.CompanyGroup)
            .filter(models.CompanyGroup.id.in_(group_ids))
            .order_by(models.CompanyGroup.name.asc())
            .all()
        )
        if len(groups) != len(group_ids):
            raise HTTPException(status_code=400, detail="Invalid group selection")

    new_user = models.User(
        username=name,
        email=employee.email,
        password_hash="",
        company_role_id=company_role.id,
        system_role_id=system_role.id,
        is_active=True,
    )
    db.add(new_user)
    db.flush()

    for group in groups:
        db.execute(
            models.user_company_groups.insert().values(
                user_id=new_user.id,
                company_group_id=group.id,
            )
        )

    db.commit()
    db.refresh(new_user)

    return {
        "id": new_user.id,
        "username": new_user.username,
        "email": new_user.email,
        "company_role_id": new_user.company_role_id,
        "role_name": company_role.name,
        "groups": [
            {
                "id": group.id,
                "name": group.name,
                "description": group.description,
                "member_count": 0,
            }
            for group in groups
        ],
        "is_active": new_user.is_active,
    }


@app.put("/people/employees/{employee_id}", response_model=schemas.PeopleEmployeeResponse)
def update_people_employee(employee_id: int, employee: schemas.PeopleEmployeeUpdate, db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.id == employee_id).first()
    if user is None:
        raise HTTPException(status_code=404, detail="Employee not found")

    name = employee.name.strip()
    role_name = employee.role.strip()
    if name == "" or role_name == "":
        raise HTTPException(status_code=400, detail="Name and role are required")

    duplicate_user = (
        db.query(models.User)
        .filter(models.User.id != employee_id)
        .filter((models.User.username == name) | (models.User.email == employee.email))
        .first()
    )
    if duplicate_user is not None:
        raise HTTPException(status_code=400, detail="Employee name or email already exists")

    company_role = get_or_create_company_role(role_name, db)
    group_ids = sorted(set(employee.group_ids))
    groups = []
    if len(group_ids) > 0:
        groups = (
            db.query(models.CompanyGroup)
            .filter(models.CompanyGroup.id.in_(group_ids))
            .order_by(models.CompanyGroup.name.asc())
            .all()
        )
        if len(groups) != len(group_ids):
            raise HTTPException(status_code=400, detail="Invalid group selection")

    user.username = name
    user.email = employee.email
    user.company_role_id = company_role.id

    db.execute(
        models.user_company_groups.delete().where(
            models.user_company_groups.c.user_id == employee_id
        )
    )
    for group in groups:
        db.execute(
            models.user_company_groups.insert().values(
                user_id=employee_id,
                company_group_id=group.id,
            )
        )

    db.commit()
    db.refresh(user)

    return {
        "id": user.id,
        "username": user.username,
        "email": user.email,
        "company_role_id": user.company_role_id,
        "role_name": company_role.name,
        "groups": [
            {
                "id": group.id,
                "name": group.name,
                "description": group.description,
                "member_count": 0,
            }
            for group in groups
        ],
        "is_active": user.is_active,
    }


@app.post("/people/groups", response_model=schemas.PeopleGroupResponse)
def create_people_group(group: schemas.PeopleGroupCreate, db: Session = Depends(get_db)):
    name = group.name.strip()
    description = group.description.strip() if group.description is not None else None
    if name == "":
        raise HTTPException(status_code=400, detail="Group name is required")

    existing_group = (
        db.query(models.CompanyGroup)
        .filter(func.lower(models.CompanyGroup.name) == name.lower())
        .first()
    )
    if existing_group is not None:
        raise HTTPException(status_code=400, detail="Group name already exists")

    new_group = models.CompanyGroup(name=name, description=description)
    db.add(new_group)
    db.commit()
    db.refresh(new_group)
    return {
        "id": new_group.id,
        "name": new_group.name,
        "description": new_group.description,
        "member_count": 0,
    }


@app.put("/people/groups/{group_id}", response_model=schemas.PeopleGroupResponse)
def update_people_group(group_id: int, group: schemas.PeopleGroupUpdate, db: Session = Depends(get_db)):
    existing_group = db.query(models.CompanyGroup).filter(models.CompanyGroup.id == group_id).first()
    if existing_group is None:
        raise HTTPException(status_code=404, detail="Group not found")

    name = group.name.strip()
    description = group.description.strip() if group.description is not None else None
    if name == "":
        raise HTTPException(status_code=400, detail="Group name is required")

    duplicate_group = (
        db.query(models.CompanyGroup)
        .filter(models.CompanyGroup.id != group_id)
        .filter(func.lower(models.CompanyGroup.name) == name.lower())
        .first()
    )
    if duplicate_group is not None:
        raise HTTPException(status_code=400, detail="Group name already exists")

    existing_group.name = name
    existing_group.description = description
    db.commit()
    db.refresh(existing_group)

    member_count = (
        db.query(models.user_company_groups)
        .filter(models.user_company_groups.c.company_group_id == group_id)
        .count()
    )
    return {
        "id": existing_group.id,
        "name": existing_group.name,
        "description": existing_group.description,
        "member_count": member_count,
    }


@app.post("/people/groups/{group_id}/members", response_model=schemas.PeopleGroupResponse)
def add_people_group_member(group_id: int, member: schemas.PeopleGroupMemberCreate, db: Session = Depends(get_db)):
    group = db.query(models.CompanyGroup).filter(models.CompanyGroup.id == group_id).first()
    if group is None:
        raise HTTPException(status_code=404, detail="Group not found")

    user = db.query(models.User).filter(models.User.id == member.user_id).first()
    if user is None:
        raise HTTPException(status_code=404, detail="Employee not found")

    existing_membership = (
        db.query(models.user_company_groups)
        .filter(models.user_company_groups.c.user_id == member.user_id)
        .filter(models.user_company_groups.c.company_group_id == group_id)
        .first()
    )
    if existing_membership is not None:
        raise HTTPException(status_code=400, detail="Employee is already in this group")

    db.execute(
        models.user_company_groups.insert().values(
            user_id=member.user_id,
            company_group_id=group_id,
        )
    )
    db.commit()

    member_count = (
        db.query(models.user_company_groups)
        .filter(models.user_company_groups.c.company_group_id == group_id)
        .count()
    )
    return {
        "id": group.id,
        "name": group.name,
        "description": group.description,
        "member_count": member_count,
    }


@app.delete(
    "/people/groups/{group_id}/members/{user_id}",
    response_model=schemas.PeopleGroupResponse,
)
def remove_people_group_member(
    group_id: int,
    user_id: int,
    db: Session = Depends(get_db),
):
    group = db.query(models.CompanyGroup).filter(models.CompanyGroup.id == group_id).first()
    if group is None:
        raise HTTPException(status_code=404, detail="Group not found")

    user = db.query(models.User).filter(models.User.id == user_id).first()
    if user is None:
        raise HTTPException(status_code=404, detail="Employee not found")

    delete_result = db.execute(
        models.user_company_groups.delete()
        .where(models.user_company_groups.c.user_id == user_id)
        .where(models.user_company_groups.c.company_group_id == group_id)
    )
    if delete_result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Employee is not in this group")

    db.commit()

    member_count = (
        db.query(models.user_company_groups)
        .filter(models.user_company_groups.c.company_group_id == group_id)
        .count()
    )
    return {
        "id": group.id,
        "name": group.name,
        "description": group.description,
        "member_count": member_count,
    }


@app.post("/people/groups/import", response_model=schemas.PeopleGroupImportResponse)
def import_people_groups(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload the provided .xlsx group template.")

    try:
        workbook = load_workbook(file.file, data_only=True)
        worksheet = workbook.active
    except Exception:
        raise HTTPException(status_code=400, detail="The uploaded Excel file could not be read.")

    headers = [worksheet.cell(1, column).value for column in range(1, len(GROUP_IMPORT_HEADERS) + 1)]
    if headers != GROUP_IMPORT_HEADERS:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid template columns. Expected: {', '.join(GROUP_IMPORT_HEADERS)}.",
        )

    errors = []
    rows = []
    seen_names = set()
    existing_names = {
        name.lower()
        for (name,) in db.query(models.CompanyGroup.name).all()
    }

    for row_number in range(2, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, 3)]
        if all(value is None or str(value).strip() == "" for value in values):
            continue

        name = str(values[0]).strip() if values[0] is not None else ""
        description = str(values[1]).strip() if values[1] is not None else ""

        if name == "":
            errors.append(f"Row {row_number}: Group Name is required.")

        name_key = name.lower()
        if name_key != "" and name_key in seen_names:
            errors.append(f"Row {row_number}: Group Name is duplicated in the file.")
        if name_key != "" and name_key in existing_names:
            errors.append(f"Row {row_number}: Group Name already exists in the database.")

        seen_names.add(name_key)
        rows.append((name, description))

    if len(rows) == 0:
        errors.append("The uploaded template does not contain any group rows.")

    if len(errors) > 0:
        return {"created_count": 0, "errors": errors}

    for name, description in rows:
        db.add(
            models.CompanyGroup(
                name=name,
                description=description if description != "" else None,
            )
        )

    db.commit()
    return {"created_count": len(rows), "errors": []}


@app.post("/people/employees/import", response_model=schemas.PeopleEmployeeImportResponse)
def import_people_employees(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload the provided .xlsx employee template.")

    try:
        workbook = load_workbook(file.file, data_only=True)
        worksheet = workbook.active
    except Exception:
        raise HTTPException(status_code=400, detail="The uploaded Excel file could not be read.")

    headers = [worksheet.cell(1, column).value for column in range(1, len(EMPLOYEE_IMPORT_HEADERS) + 1)]
    if headers != EMPLOYEE_IMPORT_HEADERS:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid template columns. Expected: {', '.join(EMPLOYEE_IMPORT_HEADERS)}.",
        )

    errors = []
    rows = []
    seen_names = set()
    seen_emails = set()
    existing_names = {
        username.lower()
        for (username,) in db.query(models.User.username).all()
    }
    existing_emails = {
        email.lower()
        for (email,) in db.query(models.User.email).all()
    }

    for row_number in range(2, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, 5)]
        if all(value is None or str(value).strip() == "" for value in values):
            continue

        name = str(values[0]).strip() if values[0] is not None else ""
        email = str(values[1]).strip() if values[1] is not None else ""
        company_role_name = str(values[2]).strip() if values[2] is not None else ""
        system_role_name = str(values[3]).strip() if values[3] is not None else ""

        if name == "":
            errors.append(f"Row {row_number}: Employee Name is required.")
        if email == "":
            errors.append(f"Row {row_number}: Employee Email Address is required.")
        else:
            try:
                email_validator.validate_python(email)
            except ValidationError:
                errors.append(f"Row {row_number}: Employee Email Address is not a valid email.")
        if company_role_name == "":
            errors.append(f"Row {row_number}: Employee Company Role is required.")
        if system_role_name == "":
            errors.append(f"Row {row_number}: Employee System Role is required.")

        name_key = name.lower()
        email_key = email.lower()
        if name_key != "" and name_key in seen_names:
            errors.append(f"Row {row_number}: Employee Name is duplicated in the file.")
        if email_key != "" and email_key in seen_emails:
            errors.append(f"Row {row_number}: Employee Email Address is duplicated in the file.")
        if name_key != "" and name_key in existing_names:
            errors.append(f"Row {row_number}: Employee Name already exists in the database.")
        if email_key != "" and email_key in existing_emails:
            errors.append(f"Row {row_number}: Employee Email Address already exists in the database.")

        seen_names.add(name_key)
        seen_emails.add(email_key)
        rows.append((name, email, company_role_name, system_role_name))

    if len(rows) == 0:
        errors.append("The uploaded template does not contain any employee rows.")

    if len(errors) > 0:
        return {"created_count": 0, "errors": errors}

    for name, email, company_role_name, system_role_name in rows:
        company_role = get_or_create_company_role(company_role_name, db)
        system_role = get_or_create_system_role(system_role_name, db)
        db.add(
            models.User(
                username=name,
                email=email,
                password_hash="",
                company_role_id=company_role.id,
                system_role_id=system_role.id,
                is_active=True,
            )
        )

    db.commit()
    return {"created_count": len(rows), "errors": []}


def get_campaign_form_ids(campaign_id: int, db: Session):
    rule_form_ids = {
        form_id
        for (form_id,) in db.query(models.CampaignEvaluationRule.form_id)
        .filter(models.CampaignEvaluationRule.campaign_id == campaign_id)
        .distinct()
        .all()
    }
    filled_form_ids = {
        form_id
        for (form_id,) in db.query(models.FilledForm.form_id)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .distinct()
        .all()
    }

    return sorted(rule_form_ids | filled_form_ids)


def get_campaign_participant_count(campaign_id: int, db: Session):
    evaluator_ids = {
        user_id
        for (user_id,) in db.query(models.FilledForm.evaluator_id)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .distinct()
        .all()
    }
    evaluatee_ids = {
        user_id
        for (user_id,) in db.query(models.FilledForm.evaluatee_id)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .distinct()
        .all()
    }

    return len(evaluator_ids | evaluatee_ids)


def get_campaign_participants(campaign_id: int, db: Session):
    assigned_group_ids = get_campaign_assigned_group_ids(campaign_id, db)
    if len(assigned_group_ids) == 0:
        return []

    rows = (
        db.query(models.User, models.CompanyRole, models.CompanyGroup)
        .join(models.CompanyRole, models.CompanyRole.id == models.User.company_role_id)
        .join(models.user_company_groups, models.user_company_groups.c.user_id == models.User.id)
        .join(models.CompanyGroup, models.CompanyGroup.id == models.user_company_groups.c.company_group_id)
        .filter(models.user_company_groups.c.company_group_id.in_(assigned_group_ids))
        .filter(models.User.is_active.is_(True))
        .order_by(models.User.username.asc(), models.CompanyGroup.name.asc())
        .all()
    )

    participants_by_id = {}
    evaluations_left_by_user_id = dict(
        db.query(models.FilledForm.evaluator_id, func.count(models.FilledForm.id))
        .filter(models.FilledForm.campaign_id == campaign_id)
        .filter(models.FilledForm.finish_date.is_(None))
        .group_by(models.FilledForm.evaluator_id)
        .all()
    )
    for user, role, group in rows:
        participant = participants_by_id.setdefault(
            user.id,
            {
                "user_id": user.id,
                "name": user.username,
                "email": user.email,
                "profile_image_url": user.profile_image_url,
                "role_name": role.name,
                "groups": [],
                "evaluations_left": evaluations_left_by_user_id.get(user.id, 0),
            },
        )
        participant["groups"].append(group.name)

    return list(participants_by_id.values())


@app.get("/users/{user_id}/assigned-evaluations", response_model=list[schemas.AssignedEvaluationResponse])
def get_assigned_evaluations(user_id: int, db: Session = Depends(get_db)):
    user = (
        db.query(models.User)
        .filter(models.User.id == user_id)
        .filter(models.User.is_active.is_(True))
        .first()
    )
    if user is None:
        raise HTTPException(status_code=404, detail="User not found")

    rows = (
        db.query(
            models.FilledForm,
            models.Campaign,
            models.Form,
            models.FormStatus,
            models.User,
        )
        .join(models.Campaign, models.Campaign.id == models.FilledForm.campaign_id)
        .join(models.Form, models.Form.id == models.FilledForm.form_id)
        .join(models.FormStatus, models.FormStatus.id == models.FilledForm.status_id)
        .join(models.User, models.User.id == models.FilledForm.evaluatee_id)
        .filter(models.FilledForm.evaluator_id == user_id)
        .order_by(
            models.FilledForm.finish_date.isnot(None).asc(),
            models.Campaign.end_date.asc().nullslast(),
            models.FilledForm.created_at.desc(),
        )
        .all()
    )

    evaluations = []
    for filled_form, campaign, form, status, evaluatee in rows:
        questions = form.questions if isinstance(form.questions, list) else []
        answers = filled_form.answers if isinstance(filled_form.answers, list) else []
        answered_count = sum(1 for answer in answers if isinstance(answer, dict) and len(answer) > 0)

        evaluations.append({
            "id": filled_form.id,
            "campaign_id": campaign.id,
            "campaign_name": campaign.name,
            "form_id": form.id,
            "form_name": form.name,
            "form_description": form.description,
            "evaluatee_id": evaluatee.id,
            "evaluatee_name": evaluatee.username,
            "status_name": status.name,
            "due_date": campaign.end_date,
            "finish_date": filled_form.finish_date,
            "question_count": len(questions),
            "answered_count": min(answered_count, len(questions)) if len(questions) > 0 else answered_count,
            "created_at": filled_form.created_at,
        })

    return evaluations


@app.get("/users/{user_id}/employee-dashboard", response_model=schemas.EmployeeDashboardResponse)
def get_employee_dashboard(user_id: int, db: Session = Depends(get_db)):
    evaluations = get_assigned_evaluations(user_id, db)
    open_evaluations = [evaluation for evaluation in evaluations if evaluation["finish_date"] is None]
    completed_evaluations = [evaluation for evaluation in evaluations if evaluation["finish_date"] is not None]
    active_campaign = (
        db.query(models.Campaign)
        .join(models.FilledForm, models.FilledForm.campaign_id == models.Campaign.id)
        .filter(models.FilledForm.evaluator_id == user_id)
        .filter(models.Campaign.is_active.is_(True))
        .order_by(models.Campaign.end_date.asc().nullslast(), models.Campaign.start_date.asc())
        .first()
    )
    next_campaign = None
    if active_campaign is None:
        next_campaign = (
            db.query(models.Campaign)
            .filter(models.Campaign.start_date >= date.today())
            .filter(models.Campaign.is_active.is_(True))
            .order_by(models.Campaign.start_date.asc())
            .first()
        )

    campaign_date_label = "Campaign ends" if active_campaign is not None else "Next campaign"
    campaign_date = (
        active_campaign.end_date
        if active_campaign is not None
        else next_campaign.start_date if next_campaign is not None else None
    )
    campaign_date_empty_text = "No deadline" if active_campaign is not None else "Not announced yet"
    latest_result = completed_evaluations[0] if len(completed_evaluations) > 0 else None

    return {
        "evaluations": evaluations,
        "open_count": len(open_evaluations),
        "completed_count": len(completed_evaluations),
        "next_review_date": campaign_date,
        "campaign_date_label": campaign_date_label,
        "campaign_date": campaign_date,
        "campaign_date_empty_text": campaign_date_empty_text,
        "latest_result": latest_result,
        "history": [
            {
                "id": evaluation["id"],
                "title": evaluation["form_name"],
                "subtitle": evaluation["campaign_name"],
                "status": "Submitted" if evaluation["finish_date"] is not None else evaluation["status_name"],
                "occurred_at": evaluation["finish_date"] or evaluation["created_at"],
            }
            for evaluation in evaluations[:3]
        ],
    }


@app.get("/dashboard", response_model=schemas.DashboardResponse)
def get_dashboard(campaign_id: int | None = Query(default=None), db: Session = Depends(get_db)):
    campaigns = (
        db.query(models.Campaign)
        .order_by(models.Campaign.is_active.desc(), models.Campaign.start_date.desc())
        .all()
    )

    selected_campaign = None
    if campaign_id is not None:
        selected_campaign = next((campaign for campaign in campaigns if campaign.id == campaign_id), None)

    if selected_campaign is None and len(campaigns) > 0:
        selected_campaign = campaigns[0]

    today = date.today()
    upcoming_campaigns = (
        db.query(models.Campaign)
        .filter(models.Campaign.start_date >= today)
        .order_by(models.Campaign.start_date.asc())
        .limit(3)
        .all()
    )

    upcoming_reviews = [
        {
            "campaign_id": campaign.id,
            "name": campaign.name,
            "start_date": campaign.start_date,
            "participant_count": get_campaign_participant_count(campaign.id, db),
            "form_count": len(get_campaign_form_ids(campaign.id, db)),
        }
        for campaign in upcoming_campaigns
    ]
    upcoming_deadlines = sorted(
        [
            {
                "campaign_id": campaign.id,
                "name": campaign.name,
                "deadline_type": "Starts",
                "date": campaign.start_date,
            }
            for campaign in campaigns
            if campaign.start_date >= today
        ]
        + [
            {
                "campaign_id": campaign.id,
                "name": campaign.name,
                "deadline_type": "Ends",
                "date": campaign.end_date,
            }
            for campaign in campaigns
            if campaign.end_date is not None and campaign.end_date >= today
        ],
        key=lambda item: (item["date"], item["deadline_type"], item["name"], item["campaign_id"]),
    )[:8]
    campaign_summaries = [
        {
            "id": campaign.id,
            "name": campaign.name,
            "is_active": campaign.is_active,
            "start_date": campaign.start_date,
            "end_date": campaign.end_date,
        }
        for campaign in campaigns
    ]

    if selected_campaign is None:
        return {
            "campaigns": campaign_summaries,
            "selected_campaign_id": None,
            "metrics": None,
            "forms": [],
            "participants": [],
            "upcoming_reviews": upcoming_reviews,
            "upcoming_deadlines": upcoming_deadlines,
        }

    assignments = (
        db.query(models.FilledForm)
        .filter(models.FilledForm.campaign_id == selected_campaign.id)
        .all()
    )
    assignment_count = len(assignments)
    submitted_count = sum(1 for assignment in assignments if assignment.finish_date is not None)
    overdue_assignments = [
        assignment
        for assignment in assignments
        if assignment.finish_date is None
        and selected_campaign.end_date is not None
        and selected_campaign.end_date < today
    ]
    form_ids = get_campaign_form_ids(selected_campaign.id, db)
    overdue_form_ids = {assignment.form_id for assignment in overdue_assignments}
    next_campaign = next((campaign for campaign in upcoming_campaigns if campaign.id != selected_campaign.id), None)

    metrics = {
        "assignment_count": assignment_count,
        "submitted_count": submitted_count,
        "completion_rate": round((submitted_count / assignment_count) * 100) if assignment_count > 0 else 0,
        "awaiting_submission_count": assignment_count - submitted_count,
        "not_started_count": assignment_count - submitted_count,
        "in_progress_count": 0,
        "overdue_count": len(overdue_assignments),
        "overdue_forms_count": len(overdue_form_ids),
        "upcoming_reviews_count": len(upcoming_reviews),
        "next_review_date": next_campaign.start_date if next_campaign is not None else None,
    }

    forms = []
    for form_id in form_ids:
        form = db.query(models.Form).filter(models.Form.id == form_id).first()
        if form is None:
            continue

        form_assignments = [assignment for assignment in assignments if assignment.form_id == form_id]
        form_submitted_count = sum(1 for assignment in form_assignments if assignment.finish_date is not None)
        form_overdue_count = sum(1 for assignment in form_assignments if assignment in overdue_assignments)
        group_names = [
            name
            for (name,) in db.query(models.CompanyGroup.name)
            .join(models.CampaignEvaluationRule, models.CampaignEvaluationRule.company_group_id == models.CompanyGroup.id)
            .filter(models.CampaignEvaluationRule.campaign_id == selected_campaign.id)
            .filter(models.CampaignEvaluationRule.form_id == form_id)
            .distinct()
            .all()
        ]

        forms.append({
            "form_id": form.id,
            "name": form.name,
            "audience": ", ".join(group_names) if len(group_names) > 0 else "No groups assigned",
            "submitted_count": form_submitted_count,
            "assignment_count": len(form_assignments),
            "due_date": selected_campaign.end_date,
            "overdue_count": form_overdue_count,
        })

    return {
        "campaigns": campaign_summaries,
        "selected_campaign_id": selected_campaign.id,
        "metrics": metrics,
        "forms": forms,
        "participants": get_campaign_participants(selected_campaign.id, db),
        "upcoming_reviews": upcoming_reviews,
        "upcoming_deadlines": upcoming_deadlines,
    }


@app.get("/campaigns", response_model=list[schemas.CampaignResponse])
def get_campaigns(db: Session = Depends(get_db)):
    campaigns = (
        db.query(models.Campaign)
        .order_by(models.Campaign.is_active.desc(), models.Campaign.start_date.desc())
        .all()
    )

    issued_counts = dict(
        db.query(models.FilledForm.campaign_id, func.count(models.FilledForm.id))
        .group_by(models.FilledForm.campaign_id)
        .all()
    )

    return [
        {
            **campaign.__dict__,
            "issued_forms_count": issued_counts.get(campaign.id, 0),
        }
        for campaign in campaigns
    ]


@app.get("/campaigns/{campaign_id}", response_model=schemas.CampaignResponse)
def get_campaign(campaign_id: int, db: Session = Depends(get_db)):
    campaign = (
        db.query(models.Campaign)
        .filter(models.Campaign.id == campaign_id)
        .first()
    )

    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    issued_forms_count = (
        db.query(models.FilledForm)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .count()
    )

    return {
        **campaign.__dict__,
        "issued_forms_count": issued_forms_count,
    }


@app.post("/campaigns", response_model=schemas.CampaignResponse)
def create_campaign(campaign: schemas.CampaignCreate, db: Session = Depends(get_db)):
    if campaign.end_date is not None and campaign.end_date < campaign.start_date:
        raise HTTPException(status_code=400, detail="End date cannot be earlier than start date")

    creator = (
        db.query(models.User)
        .filter(models.User.id == campaign.created_by)
        .filter(models.User.is_active.is_(True))
        .first()
    )

    if creator is None:
        raise HTTPException(status_code=400, detail="Invalid campaign creator")

    new_campaign = models.Campaign(
        name=campaign.name,
        description=campaign.description,
        start_date=campaign.start_date,
        end_date=campaign.end_date,
        is_active=campaign.is_active,
        comment=campaign.comment,
        created_by=campaign.created_by,
    )

    db.add(new_campaign)
    db.commit()
    db.refresh(new_campaign)

    return new_campaign


@app.put("/campaigns/{campaign_id}", response_model=schemas.CampaignResponse)
def update_campaign(campaign_id: int, campaign_update: schemas.CampaignUpdate, db: Session = Depends(get_db)):
    campaign = (
        db.query(models.Campaign)
        .filter(models.Campaign.id == campaign_id)
        .first()
    )

    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    issued_forms_count = (
        db.query(models.FilledForm)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .count()
    )

    if campaign_update.end_date is not None and campaign_update.end_date < campaign_update.start_date:
        raise HTTPException(status_code=400, detail="End date cannot be earlier than start date")

    if campaign.start_date <= date.today() and campaign_update.start_date != campaign.start_date:
        raise HTTPException(status_code=400, detail="Start date cannot be changed after the campaign has started")

    if (
        issued_forms_count > 0
        and campaign.end_date is not None
        and campaign_update.end_date is not None
        and campaign_update.end_date < campaign.end_date
    ):
        raise HTTPException(status_code=400, detail="End date cannot be moved earlier after questionnaires have been issued")

    campaign.name = campaign_update.name
    campaign.description = campaign_update.description
    campaign.start_date = campaign_update.start_date
    campaign.end_date = campaign_update.end_date
    campaign.is_active = campaign_update.is_active
    campaign.comment = campaign_update.comment

    db.commit()
    db.refresh(campaign)

    return {
        **campaign.__dict__,
        "issued_forms_count": issued_forms_count,
    }


@app.get("/campaigns/{campaign_id}/groups", response_model=schemas.CampaignGroupsResponse)
def get_campaign_groups(campaign_id: int, db: Session = Depends(get_db)):
    campaign = db.query(models.Campaign).filter(models.Campaign.id == campaign_id).first()
    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    groups = db.query(models.CompanyGroup).order_by(models.CompanyGroup.name.asc()).all()
    assigned_group_ids = [
        group_id
        for (group_id,) in db.query(models.campaign_company_groups.c.company_group_id)
        .filter(models.campaign_company_groups.c.campaign_id == campaign_id)
        .all()
    ]

    return {
        "available_groups": groups,
        "assigned_group_ids": assigned_group_ids,
    }


@app.put("/campaigns/{campaign_id}/groups", response_model=schemas.CampaignGroupsResponse)
def update_campaign_groups(campaign_id: int, payload: schemas.CampaignGroupsUpdate, db: Session = Depends(get_db)):
    campaign = db.query(models.Campaign).filter(models.Campaign.id == campaign_id).first()
    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    unique_group_ids = sorted(set(payload.group_ids))
    if len(unique_group_ids) > 0:
        matching_count = (
            db.query(models.CompanyGroup)
            .filter(models.CompanyGroup.id.in_(unique_group_ids))
            .count()
        )
        if matching_count != len(unique_group_ids):
            raise HTTPException(status_code=400, detail="One or more groups do not exist")

    db.execute(
        models.campaign_company_groups.delete().where(
            models.campaign_company_groups.c.campaign_id == campaign_id
        )
    )
    for group_id in unique_group_ids:
        db.execute(
            models.campaign_company_groups.insert().values(
                campaign_id=campaign_id,
                company_group_id=group_id,
            )
        )
    db.commit()

    groups = db.query(models.CompanyGroup).order_by(models.CompanyGroup.name.asc()).all()
    return {
        "available_groups": groups,
        "assigned_group_ids": unique_group_ids,
    }


def get_campaign_assigned_group_ids(campaign_id: int, db: Session):
    return [
        group_id
        for (group_id,) in db.query(models.campaign_company_groups.c.company_group_id)
        .filter(models.campaign_company_groups.c.campaign_id == campaign_id)
        .all()
    ]


def get_group_roles(group_id: int, db: Session):
    return (
        db.query(models.CompanyRole)
        .join(models.User, models.User.company_role_id == models.CompanyRole.id)
        .join(models.user_company_groups, models.user_company_groups.c.user_id == models.User.id)
        .filter(models.user_company_groups.c.company_group_id == group_id)
        .filter(models.User.is_active.is_(True))
        .distinct()
        .order_by(models.CompanyRole.name.asc())
        .all()
    )


def get_group_employees(group_id: int, db: Session):
    return (
        db.query(models.User, models.CompanyRole)
        .join(models.CompanyRole, models.CompanyRole.id == models.User.company_role_id)
        .join(models.user_company_groups, models.user_company_groups.c.user_id == models.User.id)
        .filter(models.user_company_groups.c.company_group_id == group_id)
        .filter(models.User.is_active.is_(True))
        .order_by(models.User.username.asc())
        .all()
    )


def build_campaign_rule_matrix(campaign_id: int, db: Session):
    assigned_group_ids = get_campaign_assigned_group_ids(campaign_id, db)
    groups = []
    if len(assigned_group_ids) > 0:
        groups = (
            db.query(models.CompanyGroup)
            .filter(models.CompanyGroup.id.in_(assigned_group_ids))
            .order_by(models.CompanyGroup.name.asc())
            .all()
        )

    existing_rules = {
        (rule.company_group_id, rule.evaluator_role_id, rule.evaluatee_role_id): rule
        for rule in db.query(models.CampaignEvaluationRule)
        .filter(models.CampaignEvaluationRule.campaign_id == campaign_id)
        .all()
    }

    matrices = []
    for group in groups:
        roles = get_group_roles(group.id, db)
        role_pairs = []
        for evaluator_role in roles:
            for evaluatee_role in roles:
                rule = existing_rules.get((group.id, evaluator_role.id, evaluatee_role.id))
                role_pairs.append({
                    "evaluator_role_id": evaluator_role.id,
                    "evaluator_role_name": evaluator_role.name,
                    "evaluatee_role_id": evaluatee_role.id,
                    "evaluatee_role_name": evaluatee_role.name,
                    "form_id": rule.form_id if rule is not None else None,
                    "rule_id": rule.id if rule is not None else None,
                })
        matrices.append({
            "group_id": group.id,
            "group_name": group.name,
            "role_pairs": role_pairs,
        })

    forms = db.query(models.Form).order_by(models.Form.name.asc()).all()
    return {
        "forms": [{"id": form.id, "name": form.name} for form in forms],
        "groups": matrices,
    }


def get_or_create_default_form_status(db: Session):
    status = (
        db.query(models.FormStatus)
        .filter(func.lower(models.FormStatus.name) == "not started")
        .first()
    )
    if status is not None:
        return status

    status = models.FormStatus(name="Not started")
    db.add(status)
    db.flush()
    return status


def get_form_id_for_employee_pair(campaign_id: int, group_id: int, evaluator, evaluatee, db: Session):
    rule = (
        db.query(models.CampaignEvaluationRule)
        .filter(models.CampaignEvaluationRule.campaign_id == campaign_id)
        .filter(models.CampaignEvaluationRule.company_group_id == group_id)
        .filter(models.CampaignEvaluationRule.evaluator_role_id == evaluator.company_role_id)
        .filter(models.CampaignEvaluationRule.evaluatee_role_id == evaluatee.company_role_id)
        .first()
    )
    return rule.form_id if rule is not None else None


def build_campaign_evaluation_matrix(campaign_id: int, db: Session):
    assigned_group_ids = get_campaign_assigned_group_ids(campaign_id, db)
    groups = []
    if len(assigned_group_ids) > 0:
        groups = (
            db.query(models.CompanyGroup)
            .filter(models.CompanyGroup.id.in_(assigned_group_ids))
            .order_by(models.CompanyGroup.name.asc())
            .all()
        )

    forms_by_pair = {
        (rule.company_group_id, rule.evaluator_role_id, rule.evaluatee_role_id): rule.form_id
        for rule in db.query(models.CampaignEvaluationRule)
        .filter(models.CampaignEvaluationRule.campaign_id == campaign_id)
        .all()
    }
    filled_forms_by_key = {
        (
            filled_form.company_group_id,
            filled_form.evaluator_id,
            filled_form.evaluatee_id,
            filled_form.form_id,
        ): filled_form
        for filled_form in db.query(models.FilledForm)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .filter(models.FilledForm.company_group_id.isnot(None))
        .all()
    }

    matrices = []
    for group in groups:
        employee_rows = get_group_employees(group.id, db)
        employees = [user for user, _role in employee_rows]
        matrices.append({
            "group_id": group.id,
            "group_name": group.name,
            "employees": [
                {
                    "id": user.id,
                    "name": user.username,
                    "role_id": role.id,
                    "role_name": role.name,
                }
                for user, role in employee_rows
            ],
            "assignments": [
                {
                    "evaluator_id": evaluator.id,
                    "evaluatee_id": evaluatee.id,
                    "form_id": form_id,
                    "filled_form_id": filled_form.id,
                    "is_completed": filled_form.finish_date is not None,
                }
                for evaluator in employees
                for evaluatee in employees
                for form_id in [forms_by_pair.get((group.id, evaluator.company_role_id, evaluatee.company_role_id))]
                if form_id is not None
                for filled_form in [filled_forms_by_key.get((group.id, evaluator.id, evaluatee.id, form_id))]
                if filled_form is not None
            ],
        })

    return {"groups": matrices}


@app.get("/campaigns/{campaign_id}/evaluation-rules", response_model=schemas.CampaignEvaluationRulesResponse)
def get_campaign_evaluation_rules(campaign_id: int, db: Session = Depends(get_db)):
    campaign = db.query(models.Campaign).filter(models.Campaign.id == campaign_id).first()
    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")
    return build_campaign_rule_matrix(campaign_id, db)


@app.put("/campaigns/{campaign_id}/evaluation-rules", response_model=schemas.CampaignEvaluationRulesResponse)
def update_campaign_evaluation_rules(campaign_id: int, payload: schemas.CampaignEvaluationRulesUpdate, db: Session = Depends(get_db)):
    campaign = db.query(models.Campaign).filter(models.Campaign.id == campaign_id).first()
    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    assigned_group_ids = set(get_campaign_assigned_group_ids(campaign_id, db))
    form_ids = {form_id for (form_id,) in db.query(models.Form.id).all()}
    valid_role_pairs_by_group = {}
    for group_id in assigned_group_ids:
        role_ids = {role.id for role in get_group_roles(group_id, db)}
        valid_role_pairs_by_group[group_id] = {
            (evaluator_role_id, evaluatee_role_id)
            for evaluator_role_id in role_ids
            for evaluatee_role_id in role_ids
        }

    seen_pairs = set()
    for rule in payload.rules:
        if rule.company_group_id not in assigned_group_ids:
            raise HTTPException(status_code=400, detail="Rule group is not assigned to this campaign")
        if rule.form_id not in form_ids:
            raise HTTPException(status_code=400, detail="Rule form does not exist")
        pair = (rule.evaluator_role_id, rule.evaluatee_role_id)
        if pair not in valid_role_pairs_by_group.get(rule.company_group_id, set()):
            raise HTTPException(status_code=400, detail="Rule role pair is not valid for the selected group")
        unique_key = (rule.company_group_id, rule.evaluator_role_id, rule.evaluatee_role_id)
        if unique_key in seen_pairs:
            raise HTTPException(status_code=400, detail="Duplicate rule for the same group and role pair")
        seen_pairs.add(unique_key)

    db.query(models.CampaignEvaluationRule).filter(
        models.CampaignEvaluationRule.campaign_id == campaign_id
    ).delete()
    for rule in payload.rules:
        db.add(models.CampaignEvaluationRule(
            campaign_id=campaign_id,
            company_group_id=rule.company_group_id,
            evaluator_role_id=rule.evaluator_role_id,
            evaluatee_role_id=rule.evaluatee_role_id,
            form_id=rule.form_id,
        ))
    db.commit()

    return build_campaign_rule_matrix(campaign_id, db)


@app.get("/campaigns/{campaign_id}/evaluation-matrix", response_model=schemas.CampaignEvaluationMatrixResponse)
def get_campaign_evaluation_matrix(campaign_id: int, db: Session = Depends(get_db)):
    campaign = db.query(models.Campaign).filter(models.Campaign.id == campaign_id).first()
    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    return build_campaign_evaluation_matrix(campaign_id, db)


@app.put("/campaigns/{campaign_id}/evaluation-matrix", response_model=schemas.CampaignEvaluationMatrixUpdateResponse)
def update_campaign_evaluation_matrix(campaign_id: int, payload: schemas.CampaignEvaluationMatrixUpdate, db: Session = Depends(get_db)):
    campaign = db.query(models.Campaign).filter(models.Campaign.id == campaign_id).first()
    if campaign is None:
        raise HTTPException(status_code=404, detail="Campaign not found")

    assigned_group_ids = set(get_campaign_assigned_group_ids(campaign_id, db))
    group_members_by_group = {
        group_id: {user.id: user for user, _role in get_group_employees(group_id, db)}
        for group_id in assigned_group_ids
    }
    scoped_group_user_pairs = {
        (group_id, evaluator_id, evaluatee_id)
        for group_id, members in group_members_by_group.items()
        for evaluator_id in members.keys()
        for evaluatee_id in members.keys()
    }
    scoped_user_pairs = {
        (evaluator_id, evaluatee_id)
        for _group_id, evaluator_id, evaluatee_id in scoped_group_user_pairs
    }

    seen_assignments = set()
    desired_assignments = set()
    for assignment in payload.assignments:
        if assignment.company_group_id not in assigned_group_ids:
            raise HTTPException(status_code=400, detail="Assignment group is not assigned to this campaign")

        group_members = group_members_by_group.get(assignment.company_group_id, {})
        evaluator = group_members.get(assignment.evaluator_id)
        evaluatee = group_members.get(assignment.evaluatee_id)
        if evaluator is None or evaluatee is None:
            raise HTTPException(status_code=400, detail="Assignment users must be active members of the selected group")

        unique_group_pair = (assignment.company_group_id, assignment.evaluator_id, assignment.evaluatee_id)
        if unique_group_pair in seen_assignments:
            raise HTTPException(status_code=400, detail="Duplicate matrix assignment")
        seen_assignments.add(unique_group_pair)

        form_id = get_form_id_for_employee_pair(
            campaign_id,
            assignment.company_group_id,
            evaluator,
            evaluatee,
            db,
        )
        if form_id is None:
            raise HTTPException(status_code=400, detail="No form rule exists for this evaluator/evaluatee role relationship")

        desired_assignments.add((assignment.company_group_id, assignment.evaluator_id, assignment.evaluatee_id, form_id))

    status = get_or_create_default_form_status(db)
    existing_forms = (
        db.query(models.FilledForm)
        .filter(models.FilledForm.campaign_id == campaign_id)
        .all()
    )
    existing_by_key = {
        (
            filled_form.company_group_id,
            filled_form.evaluator_id,
            filled_form.evaluatee_id,
            filled_form.form_id,
        ): filled_form
        for filled_form in existing_forms
        if filled_form.company_group_id is not None
    }

    created_count = 0
    for group_id, evaluator_id, evaluatee_id, form_id in desired_assignments:
        if (group_id, evaluator_id, evaluatee_id, form_id) in existing_by_key:
            continue
        db.add(models.FilledForm(
            campaign_id=campaign_id,
            company_group_id=group_id,
            evaluator_id=evaluator_id,
            evaluatee_id=evaluatee_id,
            form_id=form_id,
            status_id=status.id,
            finish_date=None,
            answers=[],
        ))
        created_count += 1

    removed_count = 0
    kept_completed_count = 0
    for filled_form in existing_forms:
        key = (
            filled_form.company_group_id,
            filled_form.evaluator_id,
            filled_form.evaluatee_id,
            filled_form.form_id,
        )
        if filled_form.company_group_id is None:
            if (
                filled_form.finish_date is None
                and (filled_form.evaluator_id, filled_form.evaluatee_id) in scoped_user_pairs
            ):
                db.delete(filled_form)
                removed_count += 1
            elif filled_form.finish_date is not None:
                kept_completed_count += 1
            continue
        if filled_form.company_group_id not in assigned_group_ids:
            continue
        if (
            filled_form.company_group_id,
            filled_form.evaluator_id,
            filled_form.evaluatee_id,
        ) not in scoped_group_user_pairs:
            continue
        if key in desired_assignments:
            continue
        if filled_form.finish_date is not None:
            kept_completed_count += 1
            continue
        db.delete(filled_form)
        removed_count += 1

    db.commit()
    matrix = build_campaign_evaluation_matrix(campaign_id, db)
    return {
        **matrix,
        "created_count": created_count,
        "removed_count": removed_count,
        "kept_completed_count": kept_completed_count,
    }


@app.get("/forms", response_model=list[schemas.FormResponse])
def get_forms(db: Session = Depends(get_db)):
    return db.query(models.Form).order_by(models.Form.id.desc()).all()


@app.get("/form-templates", response_model=list[schemas.FormTemplateResponse])
def get_form_templates(db: Session = Depends(get_db)):
    return db.query(models.FormTemplate).order_by(models.FormTemplate.id.asc()).all()


def clone_template_questions(questions: list[dict]):
    cloned_questions = deepcopy(questions)
    for question in cloned_questions:
        if isinstance(question, dict):
            question["id"] = str(uuid4())
    return cloned_questions


def validate_choice_question_options(question: dict, question_index: int):
    question_type = question.get("type")
    if question_type not in {"Single choice", "Multiple choice"}:
        return

    options = question.get("options")
    if not isinstance(options, list):
        raise HTTPException(
            status_code=400,
            detail=f"Question {question_index + 1} needs at least two answer options",
        )

    normalized_ids = []
    normalized_labels = []
    for option_index, option in enumerate(options):
        if not isinstance(option, dict):
            raise HTTPException(
                status_code=400,
                detail=f"Question {question_index + 1} option {option_index + 1} is invalid",
            )
        option_id = option.get("id")
        option_label = option.get("label")
        if not isinstance(option_id, str) or option_id.strip() == "":
            raise HTTPException(
                status_code=400,
                detail=f"Question {question_index + 1} option {option_index + 1} needs an id",
            )
        if not isinstance(option_label, str) or option_label.strip() == "":
            raise HTTPException(
                status_code=400,
                detail=f"Question {question_index + 1} option {option_index + 1} needs a label",
            )
        normalized_ids.append(option_id.strip())
        normalized_labels.append(option_label.strip().lower())

    if len(normalized_labels) < 2:
        raise HTTPException(
            status_code=400,
            detail=f"Question {question_index + 1} needs at least two answer options",
        )
    if len(set(normalized_ids)) != len(normalized_ids):
        raise HTTPException(status_code=400, detail=f"Question {question_index + 1} has duplicate option ids")
    if len(set(normalized_labels)) != len(normalized_labels):
        raise HTTPException(status_code=400, detail=f"Question {question_index + 1} has duplicate option labels")


def validate_question_definitions(questions: list[dict]):
    if not isinstance(questions, list):
        raise HTTPException(status_code=400, detail="Questions must be a list")
    for index, question in enumerate(questions):
        if not isinstance(question, dict):
            raise HTTPException(status_code=400, detail=f"Question {index + 1} is invalid")
        validate_choice_question_options(question, index)


def validate_form_like_payload(payload: schemas.FormCreate, model, db: Session, item_id: int | None = None):
    name = payload.name.strip()
    if name == "":
        raise HTTPException(status_code=400, detail="Name is required")

    existing_query = db.query(model).filter(func.lower(model.name) == name.lower())
    if item_id is not None:
        existing_query = existing_query.filter(model.id != item_id)
    if existing_query.first() is not None:
        raise HTTPException(status_code=400, detail="Name already exists")

    validate_question_definitions(payload.questions)

    return {
        "name": name,
        "description": payload.description.strip() if payload.description is not None else None,
        "questions": payload.questions,
    }


@app.post("/forms", response_model=schemas.FormResponse)
def create_form(form: schemas.FormCreate, db: Session = Depends(get_db)):
    values = validate_form_like_payload(form, models.Form, db)
    if form.source_template_id is not None:
        template = db.query(models.FormTemplate).filter(models.FormTemplate.id == form.source_template_id).first()
        if template is None:
            raise HTTPException(status_code=404, detail="Template not found")
        values["questions"] = clone_template_questions(template.questions)
    new_form = models.Form(**values)
    db.add(new_form)
    db.commit()
    db.refresh(new_form)
    return new_form


@app.post("/form-templates", response_model=schemas.FormTemplateResponse)
def create_form_template(template: schemas.FormCreate, db: Session = Depends(get_db)):
    new_template = models.FormTemplate(**validate_form_like_payload(template, models.FormTemplate, db))
    db.add(new_template)
    db.commit()
    db.refresh(new_template)
    return new_template


@app.put("/forms/{form_id}", response_model=schemas.FormResponse)
def update_form(form_id: int, form_update: schemas.FormUpdate, db: Session = Depends(get_db)):
    form = db.query(models.Form).filter(models.Form.id == form_id).first()
    if form is None:
        raise HTTPException(status_code=404, detail="Form not found")
    values = validate_form_like_payload(form_update, models.Form, db, form_id)
    form.name = values["name"]
    form.description = values["description"]
    form.questions = values["questions"]
    db.commit()
    db.refresh(form)
    return form


@app.put("/form-templates/{template_id}", response_model=schemas.FormTemplateResponse)
def update_form_template(template_id: int, template_update: schemas.FormUpdate, db: Session = Depends(get_db)):
    template = db.query(models.FormTemplate).filter(models.FormTemplate.id == template_id).first()
    if template is None:
        raise HTTPException(status_code=404, detail="Template not found")
    values = validate_form_like_payload(template_update, models.FormTemplate, db, template_id)
    template.name = values["name"]
    template.description = values["description"]
    template.questions = values["questions"]
    db.commit()
    db.refresh(template)
    return template


@app.delete("/forms/{form_id}")
def delete_form(form_id: int, db: Session = Depends(get_db)):
    form = db.query(models.Form).filter(models.Form.id == form_id).first()
    if form is None:
        raise HTTPException(status_code=404, detail="Form not found")

    db.delete(form)
    db.commit()
    return {"deleted": True}


@app.delete("/form-templates/{template_id}")
def delete_form_template(template_id: int, db: Session = Depends(get_db)):
    template = db.query(models.FormTemplate).filter(models.FormTemplate.id == template_id).first()
    if template is None:
        raise HTTPException(status_code=404, detail="Template not found")

    db.delete(template)
    db.commit()
    return {"deleted": True}
